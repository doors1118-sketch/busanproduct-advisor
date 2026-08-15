from __future__ import annotations

import argparse
import json
import re
import shutil
import sqlite3
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
import openpyxl.worksheet._reader as openpyxl_reader


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_DB = Path(
    r"C:\Users\COMTREE\Desktop\busan-local-purchase-advisor-v2\artifacts\company_db_snapshot\chatbot_company_server_20260517.db"
)

_OPENPYXL_CAST_NUMBER = openpyxl_reader._cast_number


def _safe_cast_number(value: str) -> int | float | str:
    try:
        return _OPENPYXL_CAST_NUMBER(value)
    except ValueError:
        return str(value or "").strip()


openpyxl_reader._cast_number = _safe_cast_number


def now_text() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def normalize(value: Any) -> str:
    return re.sub(r"[\s\-_·ㆍ().,/\[\]]+", "", str(value or "").strip().lower())


def digits(value: Any) -> str:
    return re.sub(r"\D+", "", str(value or ""))


def text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def int_or_zero(value: Any) -> int:
    value_text = digits(value)
    return int(value_text) if value_text else 0


def find_first(patterns: Iterable[str], base: Path) -> Path | None:
    for pattern in patterns:
        matches = sorted(base.glob(pattern), key=lambda p: p.stat().st_mtime, reverse=True)
        if matches:
            return matches[0]
    return None


def iter_xlsx_dicts(path: Path, *, required_headers: set[str], max_scan_rows: int = 30) -> Iterable[dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        header_row: int | None = None
        headers: list[str] = []
        for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=max_scan_rows, values_only=True), start=1):
            values = [text(v) for v in row]
            if required_headers.issubset(set(values)):
                header_row = row_idx
                headers = values
                break
        if header_row is None:
            raise ValueError(f"header row not found in {path.name}: required={sorted(required_headers)}")

        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            values = [text(v) for v in row]
            if not any(values):
                continue
            record = {headers[idx]: values[idx] if idx < len(values) else "" for idx in range(len(headers)) if headers[idx]}
            yield record
    finally:
        wb.close()


def backup_db(db_path: Path) -> Path:
    backup_dir = ROOT / "artifacts" / "vendor_recommendation_quality_qa"
    backup_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = backup_dir / f"{db_path.stem}.before_vendor_enrichment_{stamp}.db"
    shutil.copy2(db_path, backup)
    return backup


def ensure_meta(conn: sqlite3.Connection) -> None:
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS vendor_recommendation_enrichment_meta (
            key TEXT PRIMARY KEY,
            value TEXT NOT NULL,
            updated_at TEXT NOT NULL
        )
        """
    )


def meta_get(conn: sqlite3.Connection, key: str) -> str | None:
    row = conn.execute("SELECT value FROM vendor_recommendation_enrichment_meta WHERE key = ?", (key,)).fetchone()
    return str(row[0]) if row else None


def meta_set(conn: sqlite3.Connection, key: str, value: str) -> None:
    conn.execute(
        """
        INSERT INTO vendor_recommendation_enrichment_meta(key, value, updated_at)
        VALUES (?, ?, ?)
        ON CONFLICT(key) DO UPDATE SET value = excluded.value, updated_at = excluded.updated_at
        """,
        (key, value, now_text()),
    )


def refresh_product_policy_summary(conn: sqlite3.Connection) -> int:
    generated_at = now_text()
    products: dict[tuple[str, str], dict[str, Any]] = {}

    def key_for(code: Any, name: Any) -> tuple[str, str] | None:
        code_text = text(code)
        name_text = text(name)
        if not code_text and not name_text:
            return None
        return (code_text, normalize(name_text) or code_text)

    def ensure(code: Any, name: Any) -> dict[str, Any] | None:
        key = key_for(code, name)
        if key is None:
            return None
        code_text, _ = key
        name_text = text(name) or code_text
        row = products.setdefault(
            key,
            {
                "detail_product_code": code_text,
                "detail_product_name": name_text,
                "is_sme_competition_product": 0,
                "is_construction_material_direct_purchase": 0,
                "required_special_note": "",
                "construction_material_note": "",
                "eligible_coop_count": 0,
                "busan_eligible_coop_count": 0,
                "busan_eligible_coops": "",
                "coop_joint_product_count": 0,
                "busan_coop_joint_product_count": 0,
                "busan_coop_joint_product_coops": "",
                "mas_suppliers": set(),
                "mall_suppliers": set(),
                "busan_suppliers": set(),
                "direct_suppliers": set(),
                "source_refs": set(),
            },
        )
        if len(text(row["detail_product_name"])) < len(name_text):
            row["detail_product_name"] = name_text
        return row

    if table_exists(conn, "procurement_product_classification"):
        for code, name in conn.execute(
            """
            SELECT classification_no, classification_name
            FROM procurement_product_classification
            WHERE IFNULL(use_yn, 'Y') = 'Y'
              AND classification_unit >= 6
              AND IFNULL(classification_name, '') <> ''
            """
        ):
            row = ensure(code, name)
            if row is not None:
                row["source_refs"].add("procurement_product_classification")

    if table_exists(conn, "ref_sme_competition_product"):
        for rec in conn.execute(
            """
            SELECT detail_category_code, detail_category_name, category_name,
                   sme_competition_target, direct_purchase_target, source_name
            FROM ref_sme_competition_product
            """
        ):
            code, detail_name, category_name, sme, direct_purchase, source_name = rec
            row = ensure(code, detail_name or category_name)
            if row is None:
                continue
            row["is_sme_competition_product"] = max(int_or_zero(sme), row["is_sme_competition_product"])
            row["is_construction_material_direct_purchase"] = max(
                int_or_zero(direct_purchase), row["is_construction_material_direct_purchase"]
            )
            row["source_refs"].add(source_name or "ref_sme_competition_product")

    if table_exists(conn, "shopping_mall_product"):
        for rec in conn.execute(
            """
            SELECT detail_product_code, detail_product_name, product_name,
                   company_internal_id, contract_status, source_name
            FROM shopping_mall_product
            """
        ):
            code, detail_name, product_name, company_id, status, source_name = rec
            row = ensure(code, detail_name or product_name)
            if row is None:
                continue
            if text(status).lower() == "active":
                row["mall_suppliers"].add(company_id)
            row["busan_suppliers"].add(company_id)
            row["source_refs"].add(source_name or "shopping_mall_product")

    if table_exists(conn, "mas_product"):
        for rec in conn.execute(
            """
            SELECT detail_product_code, detail_product_name, product_name,
                   company_internal_id, contract_status, source_name
            FROM mas_product
            """
        ):
            code, detail_name, product_name, company_id, status, source_name = rec
            row = ensure(code, detail_name or product_name)
            if row is None:
                continue
            if text(status).lower() == "active":
                row["mas_suppliers"].add(company_id)
            row["busan_suppliers"].add(company_id)
            row["source_refs"].add(source_name or "mas_product")

    if table_exists(conn, "direct_production_certificate"):
        for rec in conn.execute(
            """
            SELECT detail_product_code, detail_product_name, company_internal_id,
                   validity_status, source_name
            FROM direct_production_certificate
            """
        ):
            code, detail_name, company_id, validity, source_name = rec
            row = ensure(code, detail_name)
            if row is None:
                continue
            if text(validity).lower() == "valid":
                row["direct_suppliers"].add(company_id)
            row["busan_suppliers"].add(company_id)
            row["source_refs"].add(source_name or "direct_production_certificate")

    conn.execute("DROP TABLE IF EXISTS product_policy_summary")
    conn.execute(
        """
        CREATE TABLE product_policy_summary (
            detail_product_code TEXT,
            detail_product_name TEXT NOT NULL,
            is_sme_competition_product INTEGER NOT NULL DEFAULT 0,
            is_construction_material_direct_purchase INTEGER NOT NULL DEFAULT 0,
            required_special_note TEXT NOT NULL DEFAULT '',
            construction_material_note TEXT NOT NULL DEFAULT '',
            eligible_coop_count INTEGER NOT NULL DEFAULT 0,
            busan_eligible_coop_count INTEGER NOT NULL DEFAULT 0,
            busan_eligible_coops TEXT NOT NULL DEFAULT '',
            coop_joint_product_count INTEGER NOT NULL DEFAULT 0,
            busan_coop_joint_product_count INTEGER NOT NULL DEFAULT 0,
            busan_coop_joint_product_coops TEXT NOT NULL DEFAULT '',
            mas_active_supplier_count INTEGER NOT NULL DEFAULT 0,
            shopping_mall_active_supplier_count INTEGER NOT NULL DEFAULT 0,
            busan_company_product_count INTEGER NOT NULL DEFAULT 0,
            direct_production_valid_supplier_count INTEGER NOT NULL DEFAULT 0,
            source_refs TEXT NOT NULL DEFAULT '',
            generated_at TEXT NOT NULL
        )
        """
    )
    rows = []
    for row in products.values():
        if not row["detail_product_name"]:
            continue
        rows.append(
            (
                row["detail_product_code"],
                row["detail_product_name"],
                int(row["is_sme_competition_product"]),
                int(row["is_construction_material_direct_purchase"]),
                row["required_special_note"],
                row["construction_material_note"],
                row["eligible_coop_count"],
                row["busan_eligible_coop_count"],
                row["busan_eligible_coops"],
                row["coop_joint_product_count"],
                row["busan_coop_joint_product_count"],
                row["busan_coop_joint_product_coops"],
                len(row["mas_suppliers"]),
                len(row["mall_suppliers"]),
                len(row["busan_suppliers"]),
                len(row["direct_suppliers"]),
                "|".join(sorted(row["source_refs"])),
                generated_at,
            )
        )
    conn.executemany(
        """
        INSERT INTO product_policy_summary VALUES (
            ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?
        )
        """,
        rows,
    )
    conn.execute("CREATE INDEX idx_product_policy_summary_name ON product_policy_summary(detail_product_name)")
    conn.execute("CREATE INDEX idx_product_policy_summary_code ON product_policy_summary(detail_product_code)")
    return len(rows)


def table_exists(conn: sqlite3.Connection, name: str) -> bool:
    return (
        conn.execute(
            "SELECT 1 FROM sqlite_master WHERE name = ? AND type IN ('table', 'view')",
            (name,),
        ).fetchone()
        is not None
    )


def refresh_construction_capacity(conn: sqlite3.Connection, xlsx_path: Path | None) -> int:
    conn.execute("DROP TABLE IF EXISTS construction_capacity_license")
    conn.execute(
        """
        CREATE TABLE construction_capacity_license (
            business_no TEXT NOT NULL,
            company_name TEXT NOT NULL,
            location TEXT NOT NULL,
            branch_type TEXT NOT NULL,
            license_name TEXT NOT NULL,
            representative_license_yn TEXT NOT NULL,
            construction_capacity_amount INTEGER NOT NULL DEFAULT 0,
            source_name TEXT NOT NULL,
            source_refreshed_at TEXT NOT NULL
        )
        """
    )
    if xlsx_path is None:
        return 0
    rows = []
    for rec in iter_xlsx_dicts(
        xlsx_path,
        required_headers={"업체명", "사업자등록번호", "업체소재시군구", "본사지사구분", "면허업종", "시공능력평가금액"},
    ):
        business_no = digits(rec.get("사업자등록번호"))
        location = text(rec.get("업체소재시군구"))
        branch_type = text(rec.get("본사지사구분"))
        amount = int_or_zero(rec.get("시공능력평가금액"))
        if not business_no or "부산" not in location or branch_type != "본사" or amount <= 0:
            continue
        rows.append(
            (
                business_no,
                text(rec.get("업체명")),
                location,
                branch_type,
                text(rec.get("면허업종")),
                text(rec.get("대표면허여부")),
                amount,
                xlsx_path.name,
                now_text(),
            )
        )
    conn.executemany(
        """
        INSERT INTO construction_capacity_license VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        rows,
    )
    conn.execute("CREATE INDEX idx_construction_capacity_license_bno ON construction_capacity_license(business_no)")
    conn.execute("CREATE INDEX idx_construction_capacity_license_name ON construction_capacity_license(license_name)")
    return len(rows)


def refresh_venture_nara_order(conn: sqlite3.Connection, xlsx_path: Path | None) -> int:
    conn.execute("DROP TABLE IF EXISTS venture_nara_order")
    conn.execute(
        """
        CREATE TABLE venture_nara_order (
            business_no TEXT NOT NULL,
            company_name TEXT NOT NULL,
            demand_agency_name TEXT,
            demand_agency_region TEXT,
            product_code TEXT,
            product_name TEXT,
            detail_product_code TEXT,
            detail_product_name TEXT,
            item_identifier TEXT,
            item_name TEXT,
            venture_product_name TEXT,
            order_date TEXT,
            order_quantity REAL,
            unit_price INTEGER,
            order_amount INTEGER,
            source_name TEXT NOT NULL,
            source_refreshed_at TEXT NOT NULL
        )
        """
    )
    if xlsx_path is None:
        return 0
    rows = []
    for rec in iter_xlsx_dicts(
        xlsx_path,
        required_headers={"업체사업자등록번호", "업체", "세부품명", "벤처물품명"},
    ):
        business_no = digits(rec.get("업체사업자등록번호"))
        if not business_no:
            continue
        rows.append(
            (
                business_no,
                text(rec.get("업체")),
                text(rec.get("수요기관")),
                text(rec.get("수요기관소재시군구")),
                text(rec.get("물품분류번호")),
                text(rec.get("품명")),
                text(rec.get("세부품명번호")),
                text(rec.get("세부품명")),
                text(rec.get("물품식별번호")),
                text(rec.get("품목명")),
                text(rec.get("벤처물품명")),
                text(rec.get("주문일자") or rec.get("계약일자") or rec.get("납품요구일자")),
                float(text(rec.get("주문수량") or rec.get("수량") or "0").replace(",", "") or 0),
                int_or_zero(rec.get("계약단가") or rec.get("단가")),
                int_or_zero(rec.get("실적금액") or rec.get("금액")),
                xlsx_path.name,
                now_text(),
            )
        )
    conn.executemany(
        """
        INSERT INTO venture_nara_order VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        rows,
    )
    conn.execute("CREATE INDEX idx_venture_nara_order_bno ON venture_nara_order(business_no)")
    conn.execute("CREATE INDEX idx_venture_nara_order_detail_name ON venture_nara_order(detail_product_name)")
    return len(rows)


def refresh_candidate_enrichment_summary(conn: sqlite3.Connection) -> int:
    conn.execute("DROP TABLE IF EXISTS candidate_enrichment_summary")
    conn.execute(
        """
        CREATE TABLE candidate_enrichment_summary (
            company_id TEXT PRIMARY KEY,
            construction_capacity_summary_raw TEXT NOT NULL DEFAULT '',
            venture_nara_product_summary_raw TEXT NOT NULL DEFAULT '',
            venture_nara_order_summary_raw TEXT NOT NULL DEFAULT '',
            source_refreshed_at TEXT NOT NULL
        )
        """
    )

    company_by_bno = {
        str(row[0]): str(row[1])
        for row in conn.execute(
            """
            SELECT canonical_business_no, company_id
            FROM company_identity
            WHERE IFNULL(canonical_business_no, '') <> ''
              AND IFNULL(company_id, '') <> ''
            """
        )
    }
    summary: dict[str, dict[str, Any]] = {}

    def bucket(company_id: str) -> dict[str, Any]:
        return summary.setdefault(
            company_id,
            {
                "capacity": [],
                "venture_product": {},
                "venture_order": {},
            },
        )

    if table_exists(conn, "construction_capacity_license"):
        for row in conn.execute(
            """
            SELECT business_no, license_name, construction_capacity_amount, source_name
            FROM construction_capacity_license
            WHERE construction_capacity_amount > 0
            ORDER BY construction_capacity_amount DESC, license_name
            """
        ):
            business_no, license_name, amount, source_name = row
            company_id = company_by_bno.get(str(business_no))
            if not company_id:
                continue
            items = bucket(company_id)["capacity"]
            if len(items) < 8:
                items.append((int_or_zero(amount), f"{text(license_name)}^^{int_or_zero(amount)}^^{text(source_name)}"))

    if table_exists(conn, "venture_nara_order"):
        for row in conn.execute(
            """
            SELECT business_no, product_name, detail_product_name, item_name, venture_product_name,
                   order_date, order_amount, source_name
            FROM venture_nara_order
            """
        ):
            business_no, product_name, detail_product_name, item_name, venture_product_name, order_date, amount, source_name = row
            company_id = company_by_bno.get(str(business_no))
            if not company_id:
                continue
            b = bucket(company_id)
            product_label = text(venture_product_name) or text(item_name) or text(detail_product_name) or text(product_name)
            detail_label = text(detail_product_name) or text(product_name) or product_label
            source = text(source_name)
            product_key = (product_label, detail_label, source)
            product_amount = int_or_zero(amount)
            existing_product_amount, existing_product_item = b["venture_product"].get(product_key, (0, ""))
            if product_amount >= existing_product_amount:
                b["venture_product"][product_key] = (
                    product_amount,
                    f"{product_label}^^{detail_label}^^{text(order_date)}^^order_history_derived^^{source}",
                )

            order_key = (detail_label, source)
            order = b["venture_order"].setdefault(order_key, {"count": 0, "amount": 0, "last": "", "source": source})
            order["count"] += 1
            order["amount"] += product_amount
            if text(order_date) > str(order["last"]):
                order["last"] = text(order_date)

    if table_exists(conn, "venture_nara_product"):
        for row in conn.execute(
            """
            SELECT bizno, venture_product_name, category_name, valid_to, price_amount, source_name
            FROM venture_nara_product
            WHERE IFNULL(bizno, '') <> ''
            """
        ):
            business_no, product_name, category_name, valid_to, amount, source_name = row
            company_id = company_by_bno.get(str(business_no))
            if not company_id:
                continue
            b = bucket(company_id)
            product_label = text(product_name)
            if not product_label:
                continue
            detail_label = text(category_name) or product_label
            source = text(source_name) or "venture_nara_product_api"
            product_key = (product_label, detail_label, source)
            product_amount = int_or_zero(amount)
            existing_product_amount, _ = b["venture_product"].get(product_key, (0, ""))
            if product_amount >= existing_product_amount:
                b["venture_product"][product_key] = (
                    product_amount,
                    f"{product_label}^^{detail_label}^^{text(valid_to)}^^product_api^^{source}",
                )

    if table_exists(conn, "venture_nara_designated_company"):
        for row in conn.execute(
            """
            SELECT bizno, product_names, confirmed_date, input_date, source_name
            FROM venture_nara_designated_company
            WHERE IFNULL(bizno, '') <> ''
            """
        ):
            business_no, product_names, confirmed_date, input_date, source_name = row
            company_id = company_by_bno.get(str(business_no))
            if not company_id:
                continue
            product_label = text(product_names)
            if not product_label:
                continue
            source = text(source_name) or "venture_nara_designated_company_api"
            date_value = text(confirmed_date) or text(input_date)
            product_key = (product_label, "벤처나라 지정업체", source)
            b = bucket(company_id)
            b["venture_product"][product_key] = (
                0,
                f"{product_label}^^벤처나라 지정업체^^{date_value}^^designated_company_api^^{source}",
            )

    rows = []
    refreshed_at = now_text()
    for company_id, data in summary.items():
        capacity_items = [
            item
            for _, item in sorted(data["capacity"], key=lambda pair: (-pair[0], pair[1]))[:8]
        ]
        venture_product_items = [
            item
            for _, item in sorted(data["venture_product"].values(), key=lambda pair: (-pair[0], pair[1]))[:8]
        ]
        venture_order_items = []
        for (detail_label, source), order in sorted(
            data["venture_order"].items(),
            key=lambda pair: (-int(pair[1]["amount"]), -int(pair[1]["count"]), pair[0][0]),
        )[:8]:
            venture_order_items.append(
                f"{detail_label}^^{int(order['count'])}^^{int(order['amount'])}^^{order['last']}^^{source}"
            )
        rows.append(
            (
                company_id,
                "|||".join(capacity_items),
                "|||".join(venture_product_items),
                "|||".join(venture_order_items),
                refreshed_at,
            )
        )

    conn.executemany(
        """
        INSERT INTO candidate_enrichment_summary VALUES (?, ?, ?, ?, ?)
        """,
        rows,
    )
    conn.execute("CREATE INDEX idx_candidate_enrichment_summary_company_id ON candidate_enrichment_summary(company_id)")
    return len(rows)


def enrich_candidate_view(conn: sqlite3.Connection) -> bool:
    columns = {row[1] for row in conn.execute("PRAGMA table_info(chatbot_company_candidate_view)").fetchall()}
    required = {
        "construction_capacity_summary_raw",
        "venture_nara_product_summary_raw",
        "venture_nara_order_summary_raw",
    }
    row = conn.execute(
        "SELECT sql FROM sqlite_master WHERE name = 'chatbot_company_candidate_view' AND type = 'view'"
    ).fetchone()
    if not row or not row[0]:
        raise RuntimeError("chatbot_company_candidate_view must be a view to enrich it safely")
    current_sql = str(row[0])
    if required.issubset(columns) and "candidate_enrichment_summary" in current_sql:
        return False
    stored_base_sql = meta_get(conn, "chatbot_company_candidate_view_base_sql")
    original_sql = stored_base_sql or str(row[0])
    if not stored_base_sql and ("candidate_enrichment_summary" in original_sql or "construction_capacity_summary_raw" in original_sql):
        raise RuntimeError("base chatbot_company_candidate_view SQL is missing; restore from backup before rebuilding")
    meta_set(conn, "chatbot_company_candidate_view_base_sql", original_sql)
    match = re.match(
        r"\s*CREATE\s+VIEW\s+chatbot_company_candidate_view\s+AS\s*(?P<body>.*)\s*$",
        original_sql,
        flags=re.IGNORECASE | re.DOTALL,
    )
    if not match:
        raise RuntimeError("cannot parse chatbot_company_candidate_view SQL")
    body = match.group("body")
    enriched_sql = f"""
CREATE VIEW chatbot_company_candidate_view AS
SELECT
    base.*,
    COALESCE(enrich.construction_capacity_summary_raw, '') AS construction_capacity_summary_raw,
    COALESCE(enrich.venture_nara_product_summary_raw, '') AS venture_nara_product_summary_raw,
    COALESCE(enrich.venture_nara_order_summary_raw, '') AS venture_nara_order_summary_raw
FROM (
{body}
) base
LEFT JOIN candidate_enrichment_summary enrich
       ON enrich.company_id = base.company_id
"""
    conn.execute("DROP VIEW chatbot_company_candidate_view")
    conn.execute(enriched_sql)
    return True


def validate(conn: sqlite3.Connection) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for name in [
        "product_policy_summary",
        "construction_capacity_license",
        "venture_nara_order",
        "chatbot_company_candidate_view",
    ]:
        result[name] = conn.execute(f"SELECT COUNT(*) FROM {name}").fetchone()[0]
    cols = {row[1] for row in conn.execute("PRAGMA table_info(chatbot_company_candidate_view)").fetchall()}
    result["candidate_view_columns_present"] = sorted(
        col
        for col in (
            "construction_capacity_summary_raw",
            "venture_nara_product_summary_raw",
            "venture_nara_order_summary_raw",
        )
        if col in cols
    )
    result["candidate_view_construction_capacity_nonempty"] = conn.execute(
        """
        SELECT COUNT(*) FROM chatbot_company_candidate_view
        WHERE IFNULL(construction_capacity_summary_raw, '') <> ''
        """
    ).fetchone()[0]
    result["candidate_view_venture_order_nonempty"] = conn.execute(
        """
        SELECT COUNT(*) FROM chatbot_company_candidate_view
        WHERE IFNULL(venture_nara_order_summary_raw, '') <> ''
        """
    ).fetchone()[0]
    result["product_policy_direct_nonempty"] = conn.execute(
        """
        SELECT COUNT(*) FROM product_policy_summary
        WHERE direct_production_valid_supplier_count > 0
        """
    ).fetchone()[0]
    result["product_policy_sme_count"] = conn.execute(
        """
        SELECT COUNT(*) FROM product_policy_summary
        WHERE is_sme_competition_product = 1
        """
    ).fetchone()[0]
    return result


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--db", type=Path, default=DEFAULT_DB)
    parser.add_argument("--source-dir", type=Path, default=Path.home() / "Desktop")
    parser.add_argument("--no-backup", action="store_true")
    args = parser.parse_args()

    db_path = args.db
    if not db_path.exists():
        raise FileNotFoundError(db_path)

    venture_file = find_first(["UI-ADOWAA-002R*.xlsx"], args.source_dir)
    capacity_file = find_first(["UI-ADOAAA-010R*부산 본사*.xlsx"], args.source_dir)

    backup = None if args.no_backup else backup_db(db_path)
    conn = sqlite3.connect(db_path)
    try:
        ensure_meta(conn)
        with conn:
            product_rows = refresh_product_policy_summary(conn)
            capacity_rows = refresh_construction_capacity(conn, capacity_file)
            venture_rows = refresh_venture_nara_order(conn, venture_file)
            enrichment_rows = refresh_candidate_enrichment_summary(conn)
            view_recreated = enrich_candidate_view(conn)
            meta_set(
                conn,
                "last_vendor_recommendation_enrichment",
                json.dumps(
                    {
                        "generated_at": now_text(),
                        "product_policy_summary_rows": product_rows,
                        "construction_capacity_license_rows": capacity_rows,
                        "venture_nara_order_rows": venture_rows,
                        "candidate_enrichment_summary_rows": enrichment_rows,
                        "candidate_view_recreated": view_recreated,
                        "capacity_file": str(capacity_file) if capacity_file else "",
                        "venture_file": str(venture_file) if venture_file else "",
                    },
                    ensure_ascii=False,
                ),
            )
        result = validate(conn)
    finally:
        conn.close()

    print(json.dumps({"db": str(db_path), "backup": str(backup) if backup else "", **result}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
